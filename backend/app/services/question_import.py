"""Parse a spreadsheet, document, presentation, web page, JSON, PDF or image into
reviewable question rows.

Three deliberate properties:

1. Parsing never writes. It returns rows plus per-row problems, and the examiner
   decides what to import. A file with fourteen good rows and two broken ones should
   let fourteen questions through - after someone has looked at the two.

2. A row that cannot be understood is reported, not skipped. An importer that quietly
   drops the rows it dislikes is worse than one that fails: the examiner publishes an
   exam three questions short and finds out from a candidate.

3. Nothing here trusts the file's own claims about correctness. Every row goes through
   the same ``validate_question`` the manual editor and the API use, so an import
   cannot produce a question the authoring form would have refused.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any

from app.db.models.enums import (
    OPTION_BEARING_TYPES,
    Difficulty,
    QuestionCategory,
    QuestionType,
)
from app.services.validators import OptionDraft, ValidationError, validate_question

#: The column names an examiner is likely to have typed. Matching is case- and
#: punctuation-insensitive, because "Question Text", "question_text" and "QUESTION"
#: are all the same column to a human and none of them is more correct.
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "body": ("question", "question text", "body", "text", "problem", "statement"),
    "question_type": ("type", "question type", "qtype", "kind"),
    "subject": ("subject", "subject code", "subject name", "course"),
    "difficulty": ("difficulty", "level"),
    "topic": ("topic", "subtopic", "unit", "chapter"),
    "category": ("category", "section", "shelf"),
    "marks": ("marks", "mark", "points", "score", "weight"),
    "negative_marks": ("negative marks", "negative", "penalty", "negative mark"),
    "correct": ("correct", "answer", "correct answer", "correct option", "key"),
    "model_answer": ("model answer", "expected answer", "reference answer", "solution"),
    "explanation": ("explanation", "rationale", "why"),
    "tags": ("tags", "labels"),
    "min_words": ("min words", "minimum words"),
    "max_words": ("max words", "maximum words", "word limit"),
    "tolerance": ("tolerance", "margin"),
    "unit": ("unit", "units"),
    "case_sensitive": ("case sensitive", "case"),
}

#: Option columns are positional: "Option A".."Option F", or "A".."F".
OPTION_PATTERN = re.compile(r"^(?:option[\s_-]*)?([a-f])$", re.IGNORECASE)

#: What an examiner might write in a "type" cell for each real type.
TYPE_ALIASES: dict[str, QuestionType] = {
    "mcq": QuestionType.MCQ,
    "single": QuestionType.MCQ,
    "single choice": QuestionType.MCQ,
    "single select": QuestionType.MCQ,
    "multiple choice": QuestionType.MCQ,
    "multi select": QuestionType.MULTI_SELECT,
    "multiple select": QuestionType.MULTI_SELECT,
    "multi": QuestionType.MULTI_SELECT,
    "multiselect": QuestionType.MULTI_SELECT,
    "multiple answer": QuestionType.MULTI_SELECT,
    "multiple answers": QuestionType.MULTI_SELECT,
    "true false": QuestionType.TRUE_FALSE,
    "true/false": QuestionType.TRUE_FALSE,
    "truefalse": QuestionType.TRUE_FALSE,
    "tf": QuestionType.TRUE_FALSE,
    "fill blank": QuestionType.FILL_BLANK,
    "fill in the blank": QuestionType.FILL_BLANK,
    "fillblank": QuestionType.FILL_BLANK,
    "blank": QuestionType.FILL_BLANK,
    "numerical": QuestionType.NUMERICAL,
    "numeric": QuestionType.NUMERICAL,
    "number": QuestionType.NUMERICAL,
    "short answer": QuestionType.SHORT_ANSWER,
    "short": QuestionType.SHORT_ANSWER,
    "long answer": QuestionType.LONG_ANSWER,
    "long": QuestionType.LONG_ANSWER,
    "essay": QuestionType.LONG_ANSWER,
    "descriptive": QuestionType.LONG_ANSWER,
    "handwritten": QuestionType.IMAGE_UPLOAD,
    "image upload": QuestionType.IMAGE_UPLOAD,
    "upload": QuestionType.IMAGE_UPLOAD,
    "coding": QuestionType.CODING,
    "code": QuestionType.CODING,
    "programming": QuestionType.CODING,
    "passage": QuestionType.PASSAGE,
    "comprehension": QuestionType.PASSAGE,
}


class ImportError_(Exception):
    """The file itself could not be read - wrong format, corrupt, or empty."""


@dataclass
class ParsedOption:
    text: str
    is_correct: bool


@dataclass
class ParsedRow:
    """One candidate question, as understood, plus everything wrong with it.

    ``problems`` empty means this row would be accepted by the same validation the
    manual editor applies. It is not a promise that the question is *good* - only that
    it is well-formed.
    """

    #: 1-based, counting the header, so it matches what the examiner sees in Excel.
    row_number: int
    body: str = ""
    #: The raw text of a "Subject" cell, exactly as typed - "" when the column is
    #: absent or the cell is blank, in which case the row takes the import's default
    #: subject. Resolving this to a real ``Subject`` id needs the database, so it
    #: happens one layer up in the API, not here.
    subject_text: str = ""
    #: Filled in by the API once ``subject_text`` (or the import's default) has been
    #: resolved against the database. ``None`` until then, and still ``None`` if the
    #: text named a subject that does not exist - see ``problems`` for that case.
    subject_id: Any | None = None
    subject_name: str | None = None
    question_type: QuestionType = QuestionType.MCQ
    difficulty: Difficulty = Difficulty.MEDIUM
    category: QuestionCategory = QuestionCategory.ACADEMIC
    topic: str | None = None
    marks: float = 1.0
    negative_marks: float = 0.0
    model_answer: str | None = None
    explanation: str | None = None
    tags: list[str] = field(default_factory=list)
    min_words: int | None = None
    max_words: int | None = None
    options: list[ParsedOption] = field(default_factory=list)
    spec: dict[str, Any] | None = None
    problems: list[str] = field(default_factory=list)
    #: Set when an identical body appears earlier in the same file, or already in
    #: the bank. A duplicate is a warning an examiner should see, not a hard error -
    #: two exams legitimately reuse a question, and only they can tell.
    duplicate_of: str | None = None

    @property
    def ok(self) -> bool:
        return not self.problems

    def as_payload(self, subject_id: str) -> dict[str, Any]:
        """The row as a ``QuestionCreate`` body."""
        return {
            "subject_id": subject_id,
            "question_type": self.question_type.value,
            "difficulty": self.difficulty.value,
            "category": self.category.value,
            "topic": self.topic,
            "body": self.body,
            "marks": self.marks,
            "negative_marks": self.negative_marks,
            "model_answer": self.model_answer,
            "explanation": self.explanation,
            "tags": self.tags or None,
            "min_words": self.min_words,
            "max_words": self.max_words,
            "spec": self.spec,
            "options": [
                {"text": o.text, "is_correct": o.is_correct, "order_index": i}
                for i, o in enumerate(self.options)
            ],
        }


def _normalise_header(value: str) -> str:
    return re.sub(r"[\s_\-\.]+", " ", (value or "").strip().lower())


def _map_columns(headers: list[str]) -> tuple[dict[str, int], dict[str, int]]:
    """Return (field name -> column index, option letter -> column index)."""
    fields: dict[str, int] = {}
    options: dict[str, int] = {}
    for index, raw in enumerate(headers):
        header = _normalise_header(raw)
        if not header:
            continue
        letter_match = OPTION_PATTERN.match(header)
        if letter_match:
            options[letter_match.group(1).upper()] = index
            continue
        for name, aliases in COLUMN_ALIASES.items():
            if header in aliases and name not in fields:
                fields[name] = index
                break
    return fields, options


def _cell(row: list[str], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    value = row[index]
    return "" if value is None else str(value).strip()


def _parse_float(text: str, default: float) -> float | None:
    """None signals "present but unreadable", which is a problem worth reporting."""
    if not text:
        return default
    try:
        return float(text)
    except ValueError:
        return None


def _parse_type(text: str) -> QuestionType | None:
    if not text:
        return None
    key = _normalise_header(text)
    if key in TYPE_ALIASES:
        return TYPE_ALIASES[key]
    try:
        return QuestionType(key.replace(" ", "_"))
    except ValueError:
        return None


def _parse_correct_letters(text: str) -> list[str]:
    """ "B", "b", "A,C", "A and C", "2" all name options.

    Numbers are accepted because plenty of question banks are written with 1-based
    option numbers, and rejecting those would fail files that are perfectly clear.
    """
    found: list[str] = []
    for token in re.split(r"[,;/&]|\band\b|\s+", text.strip()):
        token = token.strip().upper().rstrip(".)")
        if not token:
            continue
        if re.fullmatch(r"[A-F]", token):
            found.append(token)
        elif re.fullmatch(r"[1-6]", token):
            found.append(chr(ord("A") + int(token) - 1))
    return list(dict.fromkeys(found))


def _match_answer_letters(answer_text: str, cells: list[str]) -> list[str]:
    """Letters an answer names, including by quoting the option's own text.

    A document that writes "Answer: Python" rather than "Answer: A" is naming the
    option, not the letter - matched here against the option cells so it resolves the
    same way a lettered answer would.
    """
    letters = _parse_correct_letters(answer_text)
    if letters or not answer_text:
        return letters
    found: list[str] = []
    for token in re.split(r"[,;]|\band\b", answer_text):
        token = token.strip().casefold()
        if not token:
            continue
        for index, cell in enumerate(cells):
            if cell and cell.strip().casefold() == token:
                letter = chr(65 + index)
                if letter not in found:
                    found.append(letter)
                break
    return found


def _row_to_question(
    row_number: int,
    row: list[str],
    fields: dict[str, int],
    options: dict[str, int],
) -> ParsedRow:
    parsed = ParsedRow(row_number=row_number)

    parsed.body = _cell(row, fields.get("body"))
    if not parsed.body:
        parsed.problems.append("The question text is empty.")

    parsed.subject_text = _cell(row, fields.get("subject"))

    type_text = _cell(row, fields.get("question_type"))
    question_type = _parse_type(type_text)
    if question_type is None:
        if type_text:
            parsed.problems.append(f"'{type_text}' is not a question type we recognise.")
        else:
            parsed.problems.append("The question type is missing.")
    else:
        parsed.question_type = question_type

    difficulty_text = _normalise_header(_cell(row, fields.get("difficulty")))
    if difficulty_text:
        try:
            parsed.difficulty = Difficulty(difficulty_text)
        except ValueError:
            parsed.problems.append(f"'{difficulty_text}' is not easy, medium or hard.")

    category_text = _normalise_header(_cell(row, fields.get("category"))).replace(" ", "_")
    if category_text:
        try:
            parsed.category = QuestionCategory(category_text)
        except ValueError:
            parsed.problems.append(f"'{category_text}' is not a question category.")

    parsed.topic = _cell(row, fields.get("topic")) or None
    parsed.model_answer = _cell(row, fields.get("model_answer")) or None
    parsed.explanation = _cell(row, fields.get("explanation")) or None
    tags_text = _cell(row, fields.get("tags"))
    parsed.tags = [t.strip() for t in re.split(r"[,;]", tags_text) if t.strip()]

    marks = _parse_float(_cell(row, fields.get("marks")), 1.0)
    if marks is None:
        parsed.problems.append("The marks value is not a number.")
    else:
        parsed.marks = marks

    negative = _parse_float(_cell(row, fields.get("negative_marks")), 0.0)
    if negative is None:
        parsed.problems.append("The negative marks value is not a number.")
    else:
        parsed.negative_marks = abs(negative)

    for name in ("min_words", "max_words"):
        text = _cell(row, fields.get(name))
        if text:
            try:
                setattr(parsed, name, int(float(text)))
            except ValueError:
                parsed.problems.append(f"The {name.replace('_', ' ')} value is not a number.")

    # --- options and the key
    correct_text = _cell(row, fields.get("correct"))
    if parsed.question_type in OPTION_BEARING_TYPES:
        if parsed.question_type is QuestionType.TRUE_FALSE:
            answer = _normalise_header(correct_text)
            truthy = answer in {"true", "t", "yes", "y", "1"}
            falsy = answer in {"false", "f", "no", "n", "0"}
            if not truthy and not falsy:
                parsed.problems.append("A true/false question needs TRUE or FALSE as its answer.")
            parsed.options = [
                ParsedOption("True", truthy),
                ParsedOption("False", falsy),
            ]
        else:
            letters = _parse_correct_letters(correct_text)
            for letter in sorted(options):
                text = _cell(row, options[letter])
                if text:
                    parsed.options.append(ParsedOption(text, letter in letters))
            if len(parsed.options) < 2:
                parsed.problems.append("Fewer than two options were found.")
            if not correct_text:
                parsed.problems.append("No correct answer is given.")
            elif not letters:
                parsed.problems.append(f"'{correct_text}' does not name any of the options.")
            elif not any(o.is_correct for o in parsed.options):
                parsed.problems.append(
                    f"The answer '{correct_text}' points at an option that is empty."
                )
            if parsed.question_type is QuestionType.MCQ:
                marked = sum(1 for o in parsed.options if o.is_correct)
                if marked > 1:
                    parsed.problems.append(
                        "A single-choice question has more than one option marked correct."
                    )
    elif parsed.question_type is QuestionType.NUMERICAL:
        answer = _parse_float(correct_text, None)  # type: ignore[arg-type]
        if answer is None:
            parsed.problems.append("A numerical question needs a numeric answer.")
        else:
            tolerance = _parse_float(_cell(row, fields.get("tolerance")), 0.0) or 0.0
            parsed.spec = {
                "answer": answer,
                "tolerance": tolerance,
                "unit": _cell(row, fields.get("unit")) or None,
            }
    elif parsed.question_type is QuestionType.FILL_BLANK:
        accepted = [a.strip() for a in re.split(r"[|;\n]", correct_text) if a.strip()]
        if not accepted:
            parsed.problems.append("A fill-in-the-blank question needs at least one answer.")
        else:
            case_text = _normalise_header(_cell(row, fields.get("case_sensitive")))
            parsed.spec = {
                "accepted_answers": accepted,
                "case_sensitive": case_text in {"yes", "y", "true", "1", "sensitive"},
            }
    elif parsed.question_type in (QuestionType.SHORT_ANSWER, QuestionType.LONG_ANSWER):
        # The "correct" column is where an examiner writing a spreadsheet naturally
        # puts the model answer for a written question.
        if not parsed.model_answer:
            parsed.model_answer = correct_text or None
        if not parsed.model_answer:
            parsed.problems.append("A written question needs a model answer to grade against.")

    # --- the same rules the authoring form and the API apply
    if parsed.ok:
        try:
            parsed.spec = validate_question(
                question_type=parsed.question_type,
                body=parsed.body,
                marks=parsed.marks,
                negative_marks=parsed.negative_marks,
                options=[
                    OptionDraft(o.text, o.is_correct, i) for i, o in enumerate(parsed.options)
                ],
                model_answer=parsed.model_answer,
                min_words=parsed.min_words,
                max_words=parsed.max_words,
                spec=parsed.spec,
                image_key=None,
            )
        except ValidationError as exc:
            parsed.problems.append(str(exc))

    return parsed


def _rows_from_csv(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 decodes any byte string
        raise ImportError_("That file is not text we can read.")

    sample = text[:4096]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel  # a single-column file sniffs as nothing; comma is fine

    try:
        # A blank row is kept, not dropped. ``parse_questions`` skips it when it walks
        # the rows, but it does so by position - drop it here instead and every row
        # number reported after it would be off by one from what the examiner sees in
        # their own spreadsheet.
        return [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    except csv.Error as exc:
        # latin-1 decodes any byte string, so a binary file that is not really CSV at
        # all - a renamed image, a corrupted upload - sails past the decode step above
        # and only trips here, mid-read, on a stray control character the reader
        # cannot make sense of as a field. That is a bad file, not a server fault.
        raise ImportError_("That file could not be read as CSV.") from exc


def xlsx_sheet_names(data: bytes) -> list[str]:
    """Every worksheet name in the workbook, in order. Empty on anything unreadable."""
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - dependency is declared
        return []
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return []
    names = list(workbook.sheetnames)
    workbook.close()
    return names


def xlsx_active_sheet_name(data: bytes) -> str | None:
    """Which sheet gets read when the examiner has not chosen one.

    Not necessarily ``sheetnames[0]`` - a workbook remembers whichever tab was on
    screen when it was last saved, which is what ``workbook.active`` reports.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - dependency is declared
        return None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return None
    sheet = workbook.active
    name = sheet.title if sheet is not None else None
    workbook.close()
    return name


def _rows_from_xlsx(data: bytes, sheet_name: str | None = None) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ImportError_("Spreadsheet support is not installed on the server.") from exc

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError_("That file could not be opened as a spreadsheet.") from exc

    # A workbook with several worksheets is common for a bank organised one subject per
    # tab. Without a name, the active sheet is used - whichever was on screen when the
    # file was saved - same as before this could be picked explicitly.
    if sheet_name is not None:
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ImportError_(f"This workbook has no sheet named '{sheet_name}'.")
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.active
    if sheet is None:
        workbook.close()
        raise ImportError_("That workbook has no sheets.")
    rows = [
        ["" if cell is None else str(cell) for cell in row]
        for row in sheet.iter_rows(values_only=True)
    ]
    workbook.close()
    # A blank row is kept, not dropped, so the row number reported for everything
    # after it still matches the row the examiner sees in Excel - the same reasoning
    # as the CSV reader just above.
    return rows


def _same_header(a: list[str], b: list[str]) -> bool:
    return [_normalise_header(x) for x in a] == [_normalise_header(x) for x in b]


def _rows_from_docx(data: bytes) -> list[list[str]]:
    """Read every table in a Word document, or fall back to numbered paragraphs.

    Examiners write Word question papers both ways, and a table is unambiguous where
    prose is not, so tables win whenever the document has any. A question bank split
    into several sections often comes as one table per section (Single Choice, then
    True/False, then Multiple Choice) - reading only the first would silently drop
    every section after it, so every table's data rows are collected, and a repeated
    header row in a later table is recognised and skipped rather than imported as a
    question.
    """
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ImportError_("Word support is not installed on the server.") from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ImportError_("That file could not be opened as a Word document.") from exc

    header: list[str] | None = None
    data_rows: list[list[str]] = []
    for table in document.tables:
        table_rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(table_rows) < 2:
            continue
        if header is None:
            header = table_rows[0]
            data_rows.extend(table_rows[1:])
        elif _same_header(table_rows[0], header):
            data_rows.extend(table_rows[1:])
        else:
            data_rows.extend(table_rows)

    if header is not None:
        return [header, *data_rows]

    return _rows_from_prose([p.text for p in document.paragraphs])


#: Bullet/checkbox glyphs examiners paste in front of a line - stripped before any of
#: the patterns below try to match, so "☑ A. Data collection" is read the same as
#: "A. Data collection".
_BULLET_PREFIX = re.compile(r"^[\s•●○◦‣▪✓✔☑☒☐➤►\-–—\*]+")

#: "1. What is X?" / "Q3) What is X?" - the numbering examiners actually type.
QUESTION_START = re.compile(r"^\s*(?:q(?:uestion)?\s*)?(\d{1,3})\s*[\.\):]\s*(.+)$", re.IGNORECASE)
#: "Question 1 (Easy)" / "Q.4 [Hard]" / "Question 7" - a labelled heading, often with the
#: body on the *next* line and the difficulty in brackets. The "Q"/"Question" prefix is
#: required here, since without the punctuation QUESTION_START demands, a bare number
#: at the start of a wrapped line would otherwise read as a new question.
QUESTION_LABEL = re.compile(
    r"^\s*q(?:uestion)?\.?\s*(?:no\.?\s*)?(\d{1,3})\s*"
    r"(?:[\(\[]\s*(easy|medium|hard)\s*[\)\]])?\s*[\.\):\-–—]?\s*(.*)$",
    re.IGNORECASE,
)
#: "A. Learning from data" / "(b) Removing data" / "C) ..." - optionally starred correct.
#: The separator is required: without it, a wrapped body line such as "BY dept_id;" or
#: "failure?" reads as option B or F. "e.g." is excluded for the same reason.
OPTION_LINE = re.compile(
    r"^\s*(?:\(([a-fA-F])\)|([a-fA-F])[\.\):])(?![A-Za-z]\.)\s*(.+?)\s*(\*)?\s*$"
)
#: "B. SELECT" / "(C) PRIMARY KEY" on an answer line - the letter, then the option's own
#: text repeated. Only the letter is the key; the text must not be mined for more
#: letters ("B. It references a column" is one answer, not B and A).
_LETTERED_ANSWER = re.compile(r"^\(?([A-F])\)?[\.\):]\s+\S", re.IGNORECASE)
#: A True/False option written without a letter at all - just the word on its own line,
#: which is how a True/False question is commonly laid out ("True" / "False", one per
#: line) rather than "A. True" / "B. False".
BARE_TRUE_FALSE_LINE = re.compile(r"^(true|false)\s*(\*)?$", re.IGNORECASE)
#: "Section B — True/False (Medium)" - a heading, not a question. Its type/difficulty
#: apply to every question under it until the next heading, but the heading text itself
#: must never end up inside a question body.
SECTION_HEADING = re.compile(r"^section\s+[a-z0-9]+\s*[-—–:]\s*(.+)$", re.IGNORECASE)
#: Longer phrases first - "Correct Answer: B" must not fall through the "correct"
#: alternative half-matched, leaving "Answer: B" unconsumed for OPTION_LINE to
#: mistake for an option lettered "A" through "F".
ANSWER_LINE = re.compile(
    r"^\s*(?:correct\s*answer|correct\s*option|answer|ans|correct)\s*[:\-]\s*(.+)$",
    re.IGNORECASE,
)
#: "Subject: Java" / "Marks: 2" / "Negative Marks: 0.5" - the metadata lines an examiner
#: types under a question in a Word paper. Matched before the "continuation line"
#: fallback so they are read as fields, not appended onto the question body. A line may
#: carry more than one of these separated by "|" - "Type: True/False | Difficulty:
#: Medium" is a single line, not two - so callers split on "|" before matching this.
META_LINE = re.compile(
    r"^\s*(type|subject|topic|difficulty|marks|negative\s*marks)\s*[:\-]\s*(.+)$", re.IGNORECASE
)
_META_COLUMN = {
    "type": "Type",
    "subject": "Subject",
    "topic": "Topic",
    "difficulty": "Difficulty",
    "marks": "Marks",
    "negativemarks": "Negative Marks",
}


def _rows_from_prose(lines: list[str]) -> list[list[str]]:
    """Turn a numbered question paper into header-shaped rows.

    Emitted in the same column order as the CSV template, so everything downstream -
    validation, duplicate detection, the preview - is identical whichever format the
    examiner started from.
    """
    header = [
        "Question",
        "Type",
        "Option A",
        "Option B",
        "Option C",
        "Option D",
        "Option E",
        "Option F",
        "Correct",
        "Subject",
        "Topic",
        "Difficulty",
        "Marks",
        "Negative Marks",
    ]
    rows: list[list[str]] = [header]

    body: str | None = None
    options: list[tuple[str, str, bool]] = []
    answer = ""
    meta: dict[str, str] = {}
    saw_bare_true = False
    saw_bare_false = False
    #: Set by a "Section X — <Type> (<Difficulty>)" heading and carried forward to
    #: every question under it, until the next heading changes it. A question's own
    #: explicit "Type:"/"Difficulty:" line always wins over these.
    section_type: QuestionType | None = None
    section_difficulty: str = ""
    #: The section state as it was when the *current* question started - a heading
    #: that appears between a question's last line and the next question's number must
    #: not be applied retroactively to the question just finished, only to the one
    #: about to begin.
    active_section_type: QuestionType | None = None
    active_section_difficulty: str = ""
    #: "(Easy)" on a "Question 3 (Easy)" heading - applies to that question only.
    heading_difficulty: str = ""

    def flush() -> None:
        nonlocal body, options, answer, meta, saw_bare_true, saw_bare_false
        if not body:
            body, options, answer, meta = None, [], "", {}
            return
        lettered = _LETTERED_ANSWER.match(answer)
        if lettered:
            answer = lettered.group(1).upper()
        cells = [""] * 6
        starred: list[str] = []
        for letter, text, is_correct in options:
            index = ord(letter.upper()) - 65
            if 0 <= index < 6:
                cells[index] = text
                if is_correct:
                    starred.append(letter.upper())

        # How many options are actually marked correct - from an inline "*" on the
        # option line, from the answer line naming more than one letter ("Answer: A, B,
        # C"), or from an answer written as the options' own text ("Answer: Python").
        # Whichever the document actually used.
        matched_letters = _match_answer_letters(answer, cells)
        correct_count = len(starred) or len(matched_letters)

        # The declared "Type:" is free text an examiner typed, and free text is wrong
        # sometimes - "Multiple Choice" used loosely for a question with several correct
        # options is the case actually seen in the wild. The actual count of correct
        # options is ground truth, so it wins when the two disagree: more than one
        # correct option can never be a single-choice question, whatever the label
        # above it says.
        declared = _parse_type(meta.get("Type", "")) or active_section_type
        no_lettered_options = not any(cells)
        looks_true_false = (
            saw_bare_true
            or saw_bare_false
            or (no_lettered_options and _normalise_header(answer) in {"true", "false"})
        )
        if declared is QuestionType.TRUE_FALSE or (declared is None and looks_true_false):
            resolved = "true_false"
        elif correct_count > 1:
            resolved = "multi_select"
        elif declared is not None:
            resolved = declared.value
        else:
            resolved = "mcq" if any(cells) else "short_answer"

        correct_cell = ",".join(starred) if starred else (",".join(matched_letters) or answer)

        rows.append(
            [
                body,
                resolved,
                *cells,
                correct_cell,
                meta.get("Subject", ""),
                meta.get("Topic", ""),
                meta.get("Difficulty", "") or heading_difficulty or active_section_difficulty,
                meta.get("Marks", ""),
                meta.get("Negative Marks", ""),
            ]
        )
        body, options, answer, meta = None, [], "", {}
        saw_bare_true, saw_bare_false = False, False

    for line in lines:
        raw = line.strip()
        if not raw:
            continue
        text = _BULLET_PREFIX.sub("", raw).strip() or raw

        label = QUESTION_LABEL.match(text)
        start = QUESTION_START.match(text)
        if label or start:
            flush()
            if label:
                body = label.group(3).strip()
                heading_difficulty = (label.group(2) or "").lower()
            else:
                body = start.group(2).strip()
                heading_difficulty = ""
            active_section_type, active_section_difficulty = section_type, section_difficulty
            continue

        heading = SECTION_HEADING.match(text)
        if heading:
            descriptor = heading.group(1).strip()
            without_parens = re.sub(r"\(.*?\)", "", descriptor).strip()
            found_type = _parse_type(without_parens)
            if found_type is not None:
                section_type = found_type
            found_difficulty = re.search(r"easy|medium|hard", descriptor, re.IGNORECASE)
            if found_difficulty:
                section_difficulty = found_difficulty.group(0).lower()
            continue

        if body is None:
            continue

        # First line under a bare "Question 3 (Easy)" heading is the body, whatever it
        # starts with - "Define ..." must not be mistaken for option D.
        if body == "":
            body = raw
            continue

        # A line may bundle several "Key: value" fields separated by "|" - match each
        # segment on its own, since META_LINE's value would otherwise swallow the rest
        # of the line as one field's text.
        segments = [s.strip() for s in text.split("|")] if "|" in text else [text]
        meta_matches = [(seg, META_LINE.match(seg)) for seg in segments]
        if any(match for _, match in meta_matches):
            for _, found_meta in meta_matches:
                if not found_meta:
                    continue
                key = _META_COLUMN.get(re.sub(r"\s+", "", found_meta.group(1)).lower())
                if key:
                    meta[key] = found_meta.group(2).strip()
            continue

        found_answer = ANSWER_LINE.match(text)
        if found_answer:
            answer = found_answer.group(1).strip()
            continue

        bare_tf = BARE_TRUE_FALSE_LINE.match(text)
        if bare_tf:
            if bare_tf.group(1).lower() == "true":
                saw_bare_true = True
            else:
                saw_bare_false = True
            continue

        option = OPTION_LINE.match(text)
        if option and len(text) < 400:
            letter = option.group(1) or option.group(2)
            options.append((letter, option.group(3).strip(), bool(option.group(4))))
            continue

        # A continuation line - part of the question, wrapped. Kept from the raw,
        # un-stripped text so a body that genuinely starts with a dash is not mangled.
        body = f"{body} {raw}".strip()

    flush()
    return rows


#: A page whose text layer holds fewer characters than this is treated as scanned.
_MIN_TEXT_LAYER_CHARS = 40
#: OCR costs a second or two a page, and the import request is synchronous.
MAX_OCR_PAGES = 50


def _page_has_images(page: Any) -> bool:
    try:
        return len(page.images) > 0
    except Exception:  # noqa: BLE001 - a malformed image stream just means "unknown"
        return False


def _rows_from_pdf(data: bytes) -> list[list[str]]:
    """Read a PDF the same way as a Word paper: numbered questions, options, answers.

    Pages with a real text layer are read directly. Scanned pages - no text layer, just
    a picture of the paper - are rendered and read with OCR. If the text layer yields
    no questions but some pages carry images, those are OCR'd too, for papers where the
    questions were pasted in as screenshots under a typed title.

    A PDF that still holds no questions - a syllabus, a topic list - is refused with a
    message pointing at AI Generate, which is what reads that kind of document.
    """
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ImportError_("PDF support is not installed on the server.") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ImportError_("That file could not be opened as a PDF.") from exc

    page_text: list[str] = [(page.extract_text() or "") for page in reader.pages]
    scanned = [i for i, text in enumerate(page_text) if len(text.strip()) < _MIN_TEXT_LAYER_CHARS]
    ocr_error: str | None = None

    def ocr(indexes: list[int]) -> bool:
        nonlocal ocr_error
        indexes = indexes[:MAX_OCR_PAGES]
        if not indexes:
            return False
        from app.services.ocr import read_pdf_pages

        try:
            for index, text in read_pdf_pages(data, indexes).items():
                page_text[index] = text
        except RuntimeError as exc:
            ocr_error = str(exc)
            return False
        return True

    def parse() -> list[list[str]]:
        return _rows_from_prose([line for text in page_text for line in text.splitlines()])

    ocr(scanned)
    rows = parse()
    if len(rows) < 2:
        with_images = [
            i for i, page in enumerate(reader.pages) if i not in scanned and _page_has_images(page)
        ]
        if ocr(with_images):
            rows = parse()

    if len(rows) < 2:
        if ocr_error and scanned:
            raise ImportError_(
                "This PDF looks scanned, but text recognition (OCR) is not available on "
                "the server, so no questions could be read from it."
            )
        raise ImportError_(
            "No questions were found in this PDF. If it's a syllabus rather than a "
            "question paper, use AI Generate instead."
        )
    return rows


# --------------------------------------------------------------------------------------
# Every other document type. Each reader turns its file into either table rows (when the
# document carries a real question table) or plain lines, and lines go through the same
# numbered-paper parser as Word and PDF - so a question reads the same whichever program
# the examiner wrote it in.
# --------------------------------------------------------------------------------------

_NO_QUESTIONS = (
    "No questions were found in this {kind}. Number each question (\"1. ...\" or "
    "\"Question 1\") with lettered options (\"A. ...\") and an \"Answer:\" line - or, "
    "if it's a syllabus rather than a question paper, use AI Generate instead."
)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "utf-8", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        # utf-16 "succeeds" on most even-length byte strings; only trust it with a BOM.
        if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        return text
    raise ImportError_("That file is not text we can read.")  # pragma: no cover


def _prose_or_fail(lines: list[str], kind: str) -> list[list[str]]:
    rows = _rows_from_prose(lines)
    if len(rows) < 2:
        raise ImportError_(_NO_QUESTIONS.format(kind=kind))
    return rows


def _looks_like_a_table(text: str) -> bool:
    """A first line naming a Question column, split by a delimiter, is a CSV in disguise."""
    first = next((line for line in text.splitlines() if line.strip()), "")
    if not any(sep in first for sep in (",", "\t", ";", "|")):
        return False
    cells = re.split(r"[,\t;|]", first)
    return any(_normalise_header(c.strip().strip('"')) in COLUMN_ALIASES["body"] for c in cells)


def _rows_from_text(data: bytes) -> list[list[str]]:
    """.txt / .md: a delimited table if it has a Question header, else a numbered paper."""
    text = _decode_text(data)
    if _looks_like_a_table(text):
        return _rows_from_csv(data)
    # Markdown decoration an examiner types around a question is not part of it.
    lines = [re.sub(r"^\s{0,3}(#{1,6}\s+|>\s?)|\*\*|__", "", line) for line in text.splitlines()]
    return _prose_or_fail(lines, "text file")


def _rtf_to_text(rtf: str) -> str:
    """Plain text of an RTF document: enough for question papers, not a full renderer.

    Keeps paragraph and line breaks, decodes ``\\'hh`` and ``\\uN`` characters, and drops
    destination groups (font tables, colour tables, pictures, metadata) whose content is
    not document text.
    """
    out: list[str] = []
    skip_depth: list[bool] = []
    skipping = False
    i = 0
    ignorable = {
        "fonttbl", "colortbl", "stylesheet", "info", "pict", "object", "header", "footer",
        "listtable", "listoverridetable", "rsidtbl", "generator", "xmlnstbl", "themedata",
        "colorschememapping", "latentstyles", "datastore", "filetbl", "revtbl",
    }
    while i < len(rtf):
        ch = rtf[i]
        if ch == "{":
            skip_depth.append(skipping)
            i += 1
            if rtf.startswith("\\*", i):
                skipping = True
            continue
        if ch == "}":
            skipping = skip_depth.pop() if skip_depth else False
            i += 1
            continue
        if ch == "\\":
            match = re.match(r"\\([a-zA-Z]+)(-?\d+)? ?|\\'([0-9a-fA-F]{2})|\\(.)", rtf[i:])
            if not match:
                i += 1
                continue
            i += match.end()
            word, arg, hexcode, symbol = match.groups()
            if skipping:
                continue
            if word:
                if word in ignorable:
                    skipping = True
                elif word in ("par", "line", "row", "sect", "page"):
                    out.append("\n")
                elif word == "tab" or word == "cell":
                    out.append("\t")
                elif word == "u" and arg:
                    out.append(chr(int(arg) % 65536))
                    if i < len(rtf) and rtf[i] == "?":
                        i += 1  # the ANSI fallback character after \uN
            elif hexcode:
                out.append(bytes([int(hexcode, 16)]).decode("cp1252", errors="replace"))
            elif symbol in ("\\", "{", "}"):
                out.append(symbol)
            continue
        if not skipping and ch not in "\r\n":
            out.append(ch)
        i += 1
    return "".join(out)


def _rows_from_rtf(data: bytes) -> list[list[str]]:
    text = _rtf_to_text(data.decode("latin-1"))
    return _prose_or_fail(text.splitlines(), "RTF document")


def _rows_from_html(data: bytes) -> list[list[str]]:
    """An HTML page: its first question table if it has one, else its text by block."""
    from html.parser import HTMLParser

    blocks = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section",
              "article", "table", "ul", "ol", "pre", "blockquote"}

    class _Collector(HTMLParser):
        def __init__(self) -> None:
            super().__init__(convert_charrefs=True)
            self.lines: list[str] = [""]
            self.tables: list[list[list[str]]] = []
            self._row: list[str] | None = None
            self._cell: list[str] | None = None
            self._skip = 0

        def handle_starttag(self, tag, attrs):
            if tag in ("script", "style", "head"):
                self._skip += 1
            elif tag == "table":
                self.tables.append([])
            elif tag == "tr" and self.tables:
                self._row = []
            elif tag in ("td", "th") and self._row is not None:
                self._cell = []
            if tag in blocks:
                self.lines.append("")

        def handle_endtag(self, tag):
            if tag in ("script", "style", "head"):
                self._skip = max(0, self._skip - 1)
            elif tag in ("td", "th") and self._cell is not None and self._row is not None:
                self._row.append(" ".join("".join(self._cell).split()))
                self._cell = None
            elif tag == "tr" and self._row is not None and self.tables:
                if any(self._row):
                    self.tables[-1].append(self._row)
                self._row = None
            if tag in blocks:
                self.lines.append("")

        def handle_data(self, text):
            if self._skip:
                return
            if self._cell is not None:
                self._cell.append(text)
            self.lines[-1] += text

    collector = _Collector()
    collector.feed(_decode_text(data))
    for table in collector.tables:
        if len(table) >= 2 and "body" in _map_columns(table[0])[0]:
            return table
    return _prose_or_fail(collector.lines, "web page")


def _zip_xml(data: bytes, kind: str):
    import zipfile

    try:
        return zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ImportError_(f"That file could not be opened as a {kind}.") from exc


def _xml_paragraphs(xml: bytes, paragraph_tag: str, text_tag: str) -> list[str]:
    """Paragraph texts from an Office XML part: each paragraph's text runs, joined.

    Tags are matched on their local name, so the namespace prefix never matters.
    """
    from xml.etree import ElementTree

    def local(tag: str) -> str:
        return tag.rsplit("}", 1)[-1]

    root = ElementTree.fromstring(xml)
    return [
        "".join(node.text or "" for node in element.iter() if local(node.tag) == text_tag)
        for element in root.iter()
        if local(element.tag) == paragraph_tag
    ]


def _rows_from_odt(data: bytes) -> list[list[str]]:
    """OpenDocument text (LibreOffice / Google Docs export)."""
    archive = _zip_xml(data, "OpenDocument text file")
    try:
        xml = archive.read("content.xml")
    except KeyError as exc:
        raise ImportError_("That file is not an OpenDocument text file.") from exc
    from xml.etree import ElementTree

    root = ElementTree.fromstring(xml)
    lines: list[str] = []
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1] in ("p", "h"):
            lines.append("".join(element.itertext()))
    return _prose_or_fail(lines, "document")


def _rows_from_pptx(data: bytes) -> list[list[str]]:
    """PowerPoint: every slide's text, slide by slide, in slide order."""
    archive = _zip_xml(data, "PowerPoint file")
    slides = sorted(
        (n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
        key=lambda n: int(re.search(r"(\d+)\.xml$", n).group(1)),  # type: ignore[union-attr]
    )
    if not slides:
        raise ImportError_("That PowerPoint file has no slides.")
    lines: list[str] = []
    for name in slides:
        lines.extend(_xml_paragraphs(archive.read(name), "p", "t"))
    return _prose_or_fail(lines, "presentation")


#: Extensions read by OCR - a photo or scan of a printed question paper.
IMAGE_EXTENSIONS = ("png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff", "gif")


def _rows_from_image(data: bytes) -> list[list[str]]:
    from app.services.ocr import read_image_lines

    try:
        text = read_image_lines(data)
    except RuntimeError as exc:
        raise ImportError_(f"Text recognition (OCR) could not read this image: {exc}") from exc
    return _prose_or_fail(text.splitlines(), "image")


def _docx_images_text(data: bytes) -> str:
    """OCR of the pictures inside a Word file - a scanned paper pasted in as images."""
    from app.services.ocr import read_image_lines

    archive = _zip_xml(data, "Word document")
    media = sorted(
        n for n in archive.namelist()
        if n.startswith("word/media/") and n.rsplit(".", 1)[-1].lower() in IMAGE_EXTENSIONS
    )
    texts: list[str] = []
    for name in media[:MAX_OCR_PAGES]:
        try:
            texts.append(read_image_lines(archive.read(name)))
        except RuntimeError:
            continue
    return "\n".join(texts)


def _rows_from_docx_any(data: bytes) -> list[list[str]]:
    """A Word file's tables or text - and, when that holds no questions, its pictures."""
    rows = _rows_from_docx(data)
    if len(rows) >= 2:
        return rows
    rows = _rows_from_prose(_docx_images_text(data).splitlines())
    if len(rows) < 2:
        raise ImportError_(_NO_QUESTIONS.format(kind="Word document"))
    return rows


def _rows_from_json(data: bytes) -> list[list[str]]:
    """A JSON list of questions: ``[{"question": ..., "options": [...], "answer": ...}]``.

    Also accepted wrapped as ``{"questions": [...]}``. Keys use the same names as the
    CSV template's columns. ``options`` may be plain strings or ``{"text", "is_correct"}``
    objects; the latter mark the answer themselves.
    """
    import json

    try:
        payload = json.loads(_decode_text(data))
    except ValueError as exc:
        raise ImportError_("That file is not valid JSON.") from exc
    items = payload.get("questions") if isinstance(payload, dict) else payload
    if not isinstance(items, list) or not all(isinstance(i, dict) for i in items):
        raise ImportError_('Expected a list of questions, or {"questions": [...]}.')

    keys: list[str] = []
    for item in items:
        for key in item:
            if key not in ("options", "choices") and key not in keys:
                keys.append(key)
    header = [*keys, *(f"Option {letter}" for letter in "ABCDEF")]
    if not any(_normalise_header(k) in COLUMN_ALIASES["correct"] for k in keys):
        header.append("Correct")
    rows = [header]
    for item in items:
        row = ["" if item.get(k) is None else str(item.get(k)) for k in keys]
        options = item.get("options") or item.get("choices") or []
        texts: list[str] = []
        marked: list[str] = []
        for index, option in enumerate(options[:6]):
            if isinstance(option, dict):
                texts.append(str(option.get("text", "")))
                if option.get("is_correct") or option.get("correct"):
                    marked.append(chr(65 + index))
            else:
                texts.append(str(option))
        row += texts + [""] * (6 - len(texts))
        if len(header) > len(keys) + 6:
            row.append(",".join(marked))
        elif marked:  # an explicit answer key exists, but only fill it if it is blank
            index = next(
                i for i, k in enumerate(keys) if _normalise_header(k) in COLUMN_ALIASES["correct"]
            )
            row[index] = row[index] or ",".join(marked)
        rows.append(row)
    return rows


def _via_libreoffice(extension: str, target: str):
    """Old binary Office formats (.doc/.ppt/.xls) need LibreOffice to be read at all."""
    import shutil
    import subprocess
    import tempfile
    from pathlib import Path

    def read(data: bytes) -> list[list[str]]:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice is None:
            raise ImportError_(
                f"Old .{extension} files can't be read on this server. Save it as "
                f".{target} and import that instead."
            )
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / f"upload.{extension}"
            source.write_bytes(data)
            try:
                subprocess.run(
                    [soffice, "--headless", "--convert-to", target, "--outdir", tmp, str(source)],
                    check=True, capture_output=True, timeout=90,
                )
                converted = (Path(tmp) / f"upload.{target}").read_bytes()
            except (subprocess.SubprocessError, OSError) as exc:
                raise ImportError_(
                    f"That .{extension} file could not be converted. Save it as .{target}."
                ) from exc
        return READERS[target](converted)

    return read


#: Extension -> reader.
READERS = {
    "csv": _rows_from_csv,
    "tsv": _rows_from_csv,
    "txt": _rows_from_text,
    "text": _rows_from_text,
    "md": _rows_from_text,
    "markdown": _rows_from_text,
    "rtf": _rows_from_rtf,
    "html": _rows_from_html,
    "htm": _rows_from_html,
    "json": _rows_from_json,
    "xlsx": _rows_from_xlsx,
    "xlsm": _rows_from_xlsx,
    "docx": _rows_from_docx_any,
    "odt": _rows_from_odt,
    "pptx": _rows_from_pptx,
    "pdf": _rows_from_pdf,
    **{ext: _rows_from_image for ext in IMAGE_EXTENSIONS},
    "doc": _via_libreoffice("doc", "docx"),
    "ppt": _via_libreoffice("ppt", "pptx"),
    "xls": _via_libreoffice("xls", "xlsx"),
}

#: Shown in error messages and handed to the upload control's ``accept`` list.
SUPPORTED_EXTENSIONS = tuple(READERS)


def parse_questions(
    *, filename: str, data: bytes, sheet_name: str | None = None
) -> list[ParsedRow]:
    """Read a file into rows. Raises ``ImportError_`` only when nothing can be read.

    ``sheet_name`` is meaningful only for a workbook with more than one worksheet;
    ignored for every other format.
    """
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    reader = READERS.get(extension)
    if reader is None:
        raise ImportError_(
            f"'{extension or filename}' is not a format we can import. Supported: "
            + ", ".join(f".{e}" for e in SUPPORTED_EXTENSIONS)
        )

    rows = (
        reader(data, sheet_name)  # type: ignore[call-arg]
        if extension in ("xlsx", "xlsm")
        else reader(data)
    )
    if len(rows) < 2:
        raise ImportError_(
            "That file has a header but no question rows." if rows else "That file is empty."
        )

    fields, options = _map_columns(rows[0])
    if "body" not in fields:
        raise ImportError_(
            "No question column was found. Name one column 'Question' and try again."
        )

    parsed: list[ParsedRow] = []
    seen: dict[str, int] = {}
    for offset, row in enumerate(rows[1:], start=2):
        if not any(str(cell or "").strip() for cell in row):
            continue
        item = _row_to_question(offset, list(row), fields, options)

        # Duplicate detection within the file. Whitespace and case are not meaningful
        # differences between two copies of the same question.
        fingerprint = re.sub(r"\s+", " ", item.body.strip().casefold())
        if fingerprint:
            if fingerprint in seen:
                item.duplicate_of = f"row {seen[fingerprint]}"
            else:
                seen[fingerprint] = offset
        parsed.append(item)

    if not parsed:
        raise ImportError_("That file has a header but no question rows.")
    return parsed


def fingerprint(body: str) -> str:
    """The comparison key used for duplicate detection, exposed for the bank check."""
    return re.sub(r"\s+", " ", (body or "").strip().casefold())


#: Handed to the UI so the download link and the parser can never drift apart.
TEMPLATE_HEADER = [
    "Question",
    "Type",
    "Subject",
    "Difficulty",
    "Topic",
    "Marks",
    "Negative Marks",
    "Option A",
    "Option B",
    "Option C",
    "Option D",
    "Correct",
    "Model Answer",
    "Explanation",
    "Tags",
]

#: "Subject" is optional. Left blank, a row takes whatever subject the import screen
#: has selected - the single-subject case every earlier template covered. Filled in,
#: it names the subject this row belongs to by its exact code or name, which is what
#: lets one file seed a multi-subject exam (Aptitude + Java + Python + SQL in one
#: import) rather than one file per subject.
TEMPLATE_ROWS = [
    [
        "What is supervised learning?",
        "mcq",
        "Machine Learning",
        "medium",
        "Machine Learning",
        "2",
        "0.5",
        "Learning without data",
        "Learning from labelled data",
        "Removing training data",
        "Encrypting data",
        "B",
        "",
        "Labelled examples are what makes it supervised.",
        "unit-1",
    ],
    [
        "State two limitations of gradient descent.",
        "short_answer",
        "Machine Learning",
        "hard",
        "Optimisation",
        "5",
        "0",
        "",
        "",
        "",
        "",
        "",
        "It can stall in local minima and is sensitive to the learning rate.",
        "",
        "unit-2",
    ],
    [
        "Backpropagation computes gradients by the chain rule.",
        "true_false",
        "",
        "easy",
        "Neural Networks",
        "1",
        "0",
        "",
        "",
        "",
        "",
        "TRUE",
        "",
        "",
        "",
    ],
]


def template_csv() -> str:
    """A filled-in example file, so the first import is not a guessing game."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(TEMPLATE_HEADER)
    writer.writerows(TEMPLATE_ROWS)
    return buffer.getvalue()
